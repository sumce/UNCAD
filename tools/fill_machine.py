# -*- coding: utf-8 -*-
"""
fill_machine.py — 按机台/设备名称从 Excel 填充 AutoCAD 清单表

用法:
    python fill_machine.py <excel.xlsx>

流程:
    1. 连接正在运行的 AutoCAD，命令行提示"框选区域"（含清单表）
    2. 命令行输入机台ID（如 MPAMT09；也支持直接输入设备/回路名称如 THC）
    3. 弹出窗口列出该机台所有回路（Main1/Main2/CVCF/THC/PM...），用户选一个
    4. 自动填充表格第一个空行块：
       配电行 = 电缆型号套模板（项目名称 = 详情里的 3P200A 等 + "配电"）
       软管行 = 软管直径套模板（项目名称 = "包塑金属软管"）
       项次编码 = Excel 项目序号
    数量列不动（来源未知，留给用户手填）。

依赖: openpyxl, pywin32, tkinter（Python 自带）
"""
import re
import sys
import time
import tkinter as tk

import openpyxl
import pywintypes
import win32com.client


def com_call(fn, *args, retries=12, delay=0.4, **kwargs):
    """AutoCAD 忙时 COM 调用会被拒绝（RPC_E_CALL_REJECTED），自动重试。"""
    last = None
    for _ in range(retries):
        try:
            return fn(*args, **kwargs)
        except pywintypes.com_error as e:
            last = e
            time.sleep(delay)
    raise last

# ---- 模板（与图纸手工示例一致）----
CABLE_DESC_TEMPLATE = r'1.名称:{cable}mm²单芯电缆\P2.配线形式:穿管或桥架敷设\P3.说明:包含电缆接头；支吊架；测试等一切主辅材料'
CONDUIT_DESC_TEMPLATE = r'1.名称:{dia}mm(1-1/2")包塑金属软管(波纹管)附镀锌接头\P2.材质:镀锌金属软管和PVC包覆'

# Excel 列（Sheet1）
COL_REGION = 0   # 所属区域
COL_MID    = 1   # 机台ID
COL_NAME   = 2   # 回路名称
COL_CABLE  = 3   # 电缆型号
COL_FR     = 4   # FR
COL_DETAIL = 5   # 详情
COL_SEQ    = 6   # 项目序号
COL_DIA    = 7   # 软管直径


def connect_acad():
    for prog in ('AutoCAD.Application.24.1', 'AutoCAD.Application'):
        try:
            return win32com.client.GetActiveObject(prog)
        except Exception:
            continue
    raise RuntimeError('未找到正在运行的 AutoCAD（COM）')


def load_excel(xlsx):
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    # 机台表：表头含"机台ID"的工作表（Sheet1）；不能用 wb.active（可能指向其他表）
    for ws in wb.worksheets:
        hdr = [str(c) if c is not None else '' for c in
               next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())]
        if any('机台ID' in h for h in hdr):
            return ws
    return wb['Sheet1'] if 'Sheet1' in wb.sheetnames else wb.active


def find_rows(ws, keyword):
    """按 机台ID 精确匹配；无果则按 回路名称 包含匹配。返回行列表（跳过空行）。"""
    kw = str(keyword).strip().upper()
    by_mid, by_name = [], []
    for row in ws.iter_rows(min_row=2, values_only=True):
        mid = row[COL_MID]
        if mid is None or str(mid).strip() == '':
            continue
        if str(mid).strip().upper() == kw:
            by_mid.append(row)
        elif kw and kw in str(row[COL_NAME] or '').upper():
            by_name.append(row)
    return by_mid if by_mid else by_name


def pick_circuit(rows):
    """弹窗口选回路。返回所选行；取消返回 None。"""
    root = tk.Tk()
    root.title('选择回路')
    root.attributes('-topmost', True)
    root.geometry('900x360')
    lb = tk.Listbox(root, font=('Microsoft YaHei', 9))
    for r in rows:
        lb.insert('end',
                  f"{r[COL_MID]} | {r[COL_NAME]} | 电缆:{r[COL_CABLE]} | {r[COL_DETAIL]} | 序号:{r[COL_SEQ]} | Φ{r[COL_DIA]}")
    lb.pack(fill='both', expand=True, padx=6, pady=6)
    result = {'row': None}

    def ok():
        sel = lb.curselection()
        if sel:
            result['row'] = rows[sel[0]]
        root.destroy()

    def dbl(event):
        sel = lb.curselection()
        if sel:
            result['row'] = rows[sel[0]]
            root.destroy()

    btn = tk.Button(root, text='确定（选择一行后点此）', command=ok, font=('Microsoft YaHei', 9))
    btn.pack(pady=6)
    lb.bind('<Double-Button-1>', dbl)
    root.mainloop()
    return result['row']


def breaker_name(detail, fallback):
    """详情列 " 3P4W 3P200A" -> "3P200A配电"；无匹配用回路名称。"""
    m = re.search(r'(\d+P\d+A)', str(detail or ''))
    return m.group(1) + '配电' if m else str(fallback or '')


def fill_table(tbl, picked, start_row=2):
    cable = str(picked[COL_CABLE] or '').strip()
    dia = str(picked[COL_DIA] or '').strip()
    seq = str(picked[COL_SEQ]).strip() if picked[COL_SEQ] is not None else ''

    # 找第一个空行块（项目名称列为空）
    r = start_row
    while r < tbl.Rows and str(com_call(tbl.GetText, r, 1)).strip():
        r += 1
    if r >= tbl.Rows:
        return 0, '表格已满'

    filled = 0
    # 配电行
    com_call(tbl.SetText, r, 1, breaker_name(picked[COL_DETAIL], picked[COL_NAME]))
    com_call(tbl.SetText, r, 2, CABLE_DESC_TEMPLATE.format(cable=cable))
    com_call(tbl.SetText, r, 3, 'M')
    if seq:
        com_call(tbl.SetText, r, 5, seq)
    filled += 1

    # 软管行（下一行）
    r2 = r + 1
    if r2 < tbl.Rows:
        com_call(tbl.SetText, r2, 1, '包塑金属软管')
        com_call(tbl.SetText, r2, 2, CONDUIT_DESC_TEMPLATE.format(dia=dia if dia else '?'))
        com_call(tbl.SetText, r2, 3, 'M')
        if seq:
            com_call(tbl.SetText, r2, 5, seq)
        filled += 1
    return filled, '配电+软管 已填入第 %d/%d 行' % (r + 1, r2 + 1) if r2 < tbl.Rows else '配电已填入第 %d 行' % (r + 1)


def main():
    if len(sys.argv) < 2:
        print('用法: python fill_machine.py <excel.xlsx>')
        sys.exit(1)
    xlsx = sys.argv[1]

    ws = load_excel(xlsx)
    acad = connect_acad()
    doc = acad.ActiveDocument
    util = doc.Utility

    # 1) 框选区域找表格
    util.Prompt('\n[填充] 请框选包含清单表的区域: ')
    ss_name = 'UNC_' + str(int(time.time() * 1000))
    ss = com_call(doc.SelectionSets.Add, ss_name)
    try:
        com_call(ss.SelectOnScreen)
    except Exception as e:
        try: com_call(ss.Delete)
        except Exception: pass
        print('[填充] 已取消框选:', e)
        sys.exit(1)
    time.sleep(0.5)  # 等 AutoCAD 处理完选择，避免枚举被拒
    tables = None
    for _ in range(10):
        try:
            tables = [e for e in ss if e.ObjectName == 'AcDbTable']
            break
        except pywintypes.com_error:
            time.sleep(0.5)
    try: com_call(ss.Delete)
    except Exception: pass
    if not tables:
        print('[填充] 所选区域中没有表格')
        sys.exit(1)
    tbl = tables[0]
    print('[填充] 找到表格:', tbl.Rows, '行 x', tbl.Columns, '列')

    # 2) 输入机台/设备名称
    util.Prompt('\n[填充] 请输入机台ID 或 设备/回路名称: ')
    try:
        keyword = com_call(util.GetString, 0, '')
    except Exception as e:
        print('[填充] 已取消输入:', e)
        sys.exit(1)
    keyword = str(keyword).strip()
    if not keyword:
        print('[填充] 未输入名称')
        sys.exit(1)

    # 3) 查 Excel
    rows = find_rows(ws, keyword)
    if not rows:
        print('[填充] Excel 中未找到:', keyword)
        sys.exit(1)
    if len(rows) == 1:
        picked = rows[0]
        print('[填充] 唯一匹配:', picked[COL_MID], picked[COL_NAME])
    else:
        print('[填充] 找到 %d 行，弹出选择窗口...' % len(rows))
        picked = pick_circuit(rows)
        if picked is None:
            print('[填充] 未选择，取消')
            sys.exit(1)
        print('[填充] 已选择:', picked[COL_MID], picked[COL_NAME])

    # 4) 填充
    filled, msg = fill_table(tbl, picked)
    print('[填充] 完成：' + msg + '，写入 %d 行' % filled)


if __name__ == '__main__':
    main()
